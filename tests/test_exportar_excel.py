"""
Exportación del reporte de movimientos a Excel.

Existe porque esta ruta se rompió sin que nada avisara: el recorrido de
pantallas la salteaba (necesita datos en la sesión, que sólo aparecen después
de generar un reporte) y devolvía 500 en producción.

La causa fue pandas 3.0: el cálculo del ancho de columna hacía
`df[col].astype(str).map(len)`, y una columna que queda entera vacía llega como
float64 de NaN. Hasta pandas 2.x `astype(str)` convertía NaN a la cadena 'nan';
desde la 3.0 lo deja como float y `len()` falla. Por eso la prueba fuerza el
caso: un período donde ningún movimiento es una compra, de modo que las tres
columnas de documento queden vacías.
"""
import base64
import json
import os
import re
import sys
import tempfile
import warnings
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path

warnings.filterwarnings('ignore')

os.environ['USUARIOS_AUTORIZADOS'] = '40555666'
os.environ['DATABASE_URI'] = 'sqlite:///' + os.path.join(
    tempfile.gettempdir(), 'labcrub_test_excel.db')
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

_bd = os.environ['DATABASE_URI'].replace('sqlite:///', '')
if os.path.exists(_bd):
    os.remove(_bd)

from app import create_app                      # noqa: E402
from app.utils import keycloak_auth             # noqa: E402

app = create_app()

fallos = []


def check(nombre, condicion, extra=''):
    print(('  OK   ' if condicion else '  FALLA') + ' ' + nombre
          + (('  -> ' + str(extra)) if not condicion and extra else ''))
    if not condicion:
        fallos.append(nombre)


def b64(datos):
    return base64.urlsafe_b64encode(json.dumps(datos).encode()).rstrip(b'=').decode()


class _Resp:
    def __init__(self, cuerpo):
        self._cuerpo, self.status_code, self.ok, self.text = cuerpo, 200, True, ''

    def json(self):
        return self._cuerpo

    def raise_for_status(self):
        pass


keycloak_auth.requests.post = lambda u, data=None, **k: _Resp(
    {'access_token': b64({'alg': 'none'}) + '.'
     + b64({'realm_access': {'roles': ['app_admin']}}) + '.AAAA'})
keycloak_auth.requests.get = lambda u, **k: _Resp(
    {'preferred_username': '40555666', 'given_name': 'Bea', 'family_name': 'Jefa',
     'email': 'b@crub.edu.ar', 'sub': 'kc'})


# --- datos: movimientos que NO son compras, para dejar vacías las columnas
#     de documento y proveedor, que es lo que rompía ----------------------
from app.models.models import db, Laboratorio, Movimiento, Producto   # noqa: E402

with app.app_context():
    db.create_all()
    if not Laboratorio.query.get('LAB001'):
        db.session.add(Laboratorio(idLaboratorio='LAB001', nombre='Química',
                                   direccion='Quintral 1250', telefono='294',
                                   email='q@crub.edu.ar'))
        db.session.add(Producto(idProducto='P001', nombre='Etanol',
                                tipoProducto='droguero', unidadMedida='Lt',
                                stockMinimo=2))
        db.session.commit()
    Movimiento.query.delete()
    ahora = datetime.now()
    for i, (tipo, cantidad) in enumerate([('ingreso', 10), ('uso', 3), ('ingreso', 5)]):
        db.session.add(Movimiento(
            idMovimiento=f'MX{i}', tipoMovimiento=tipo, cantidad=cantidad,
            unidadMedida='Lt', idProducto='P001', idLaboratorio='LAB001',
            timestamp=ahora - timedelta(days=3 - i)))
    db.session.commit()


def entrar(cliente):
    html = cliente.get('/auth/login').get_data(as_text=True)
    tok = re.search(r'name="csrf_token"[^>]*value="([^"]+)"', html).group(1)
    cliente.post('/auth/login',
                 data={'usuario': '40555666', 'contrasena': 'x', 'csrf_token': tok})


print('\n--- exportar a Excel sin ninguna compra en el período ---')
with app.test_client() as c:
    entrar(c)

    # sin haber generado un reporte, exportar avisa en vez de romper
    r = c.get('/admin/reportes/movimientos/excel')
    check('sin datos redirige en vez de fallar', r.status_code == 302, r.status_code)

    html = c.get('/admin/reportes/movimientos').get_data(as_text=True)
    tok = re.search(r'name="csrf_token"[^>]*value="([^"]+)"', html).group(1)
    desde = (datetime.now() - timedelta(days=30)).strftime('%Y-%m-%d')
    hasta = (datetime.now() + timedelta(days=1)).strftime('%Y-%m-%d')
    r = c.post('/admin/reportes/movimientos', data={
        'csrf_token': tok, 'fecha_inicial': desde, 'fecha_final': hasta,
        'laboratorio': '', 'tipo_producto': '', 'control_sedronar': ''})
    check('el reporte se genera', r.status_code == 200, r.status_code)

    with c.session_transaction() as s:
        filas = s.get('reporte_data_completo', [])
    check('hay filas en la sesión', len(filas) == 3, len(filas))
    check('las columnas de documento están vacías',
          all(f['tipo_documento'] is None and f['cuit_proveedor'] is None for f in filas))

    r = c.get('/admin/reportes/movimientos/excel')
    check('la exportación responde 200', r.status_code == 200, r.status_code)
    check('es un xlsx',
          'spreadsheetml.sheet' in (r.headers.get('Content-Type') or ''),
          r.headers.get('Content-Type'))
    check('viene como adjunto',
          'attachment' in (r.headers.get('Content-Disposition') or ''),
          r.headers.get('Content-Disposition'))

    datos = r.data

    with c.session_transaction() as s:
        filas_despues = s.get('reporte_data_completo', [])
    r = c.post('/admin/reportes/movimientos', data={
        'csrf_token': tok, 'fecha_inicial': desde, 'fecha_final': hasta,
        'laboratorio': '', 'tipo_producto': '', 'control_sedronar': ''})
    # Sólo la tabla: el resto de la página tiene textos que no vienen al caso.
    pantalla = r.get_data(as_text=True)
    pantalla = pantalla[pantalla.find('<tbody'):pantalla.find('</tbody>')]

print('\n--- el archivo abre y tiene los datos ---')
import openpyxl                                  # noqa: E402

libro = openpyxl.load_workbook(BytesIO(datos))
hoja = libro['Reporte de Movimientos']
check('una fila por movimiento más el encabezado', hoja.max_row == 4, hoja.max_row)
check('doce columnas', hoja.max_column == 12, hoja.max_column)

encabezados = [c.value for c in hoja[1]]
check('primer encabezado es Fecha', encabezados[0] == 'Fecha', encabezados[0])
check('último encabezado es CUIT Proveedor',
      encabezados[-1] == 'CUIT Proveedor', encabezados[-1])

primera = [c.value for c in hoja[2]]
check('el producto sale en la fila', primera[1] == 'Etanol', primera[1])
# En la planilla el tipo va con su código: compra e ingreso son CPR, uso es
# USA. Los movimientos cargados son ingreso, uso, ingreso.
tipos = sorted(fila[4].value for fila in hoja.iter_rows(min_row=2))
check('en la planilla el tipo va con su código', tipos == ['CPR', 'CPR', 'USA'], tipos)
check('en pantalla sigue como se cargó',
      'CPR' not in pantalla and 'USA' not in pantalla
      and 'Ingreso' in pantalla and 'Uso' in pantalla)
check('y exportar no cambia los datos del reporte',
      sorted(f['tipo_movimiento'] for f in filas_despues) == ['ingreso', 'ingreso', 'uso'],
      [f['tipo_movimiento'] for f in filas_despues])
check('las celdas vacías quedan vacías, no dicen "nan"',
      all(c.value in (None, '') for c in hoja[2][9:12]),
      [c.value for c in hoja[2][9:12]])

print('\n' + ('=== TODO OK ===' if not fallos
              else '=== ' + str(len(fallos)) + ' FALLAS: ' + str(fallos) + ' ==='))
sys.exit(1 if fallos else 0)
